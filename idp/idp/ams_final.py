#!/usr/bin/env python3
"""AMS 'finalised' check for archived patients (Chris, 29.09.2026).

  ams_final.py <ams_spill.json|ams.xlsx> <mrn_or_name> [...]
  ams_final.py <ams_spill.json|ams.xlsx> --list <file with one 'MRN<TAB>Name' per line>

A patient's row-block runs between two red separator rows. Old admissions are stacked
above the current one and shaded orange when closed; the current block stays blank until
it is sealed. A patient counts as FINALISED only when their LAST block (latest sheet,
lowest block) is shaded orange (any orange: E69138, FF9900, E59344 ...) over >= 70% of its
cells in columns A:N.
Only the last entry of the block counts (old admissions are stacked above it). Match by MRN first; by normalised full name when the MRN is absent.
Prints one line per patient: FINALISED / OPEN / NOT FOUND, with sheet and rows. No patient
text beyond the name and MRN that were passed in.
"""
import sys, io, json, base64, re, colorsys
import openpyxl

MONTHS = {"jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6, "jul": 7, "aug": 8,
          "sep": 9, "oct": 10, "nov": 11, "dec": 12}


def load(p):
    if p.endswith(".json"):
        d = json.load(open(p))
        return openpyxl.load_workbook(io.BytesIO(base64.b64decode(d["content"])))
    return openpyxl.load_workbook(p)


def rgb(c):
    f = c.fill
    if not f or not f.fill_type or f.fgColor is None or f.fgColor.type != "rgb":
        return None
    v = f.fgColor.rgb or ""
    return v[-6:].upper() if len(v) >= 6 else None


def is_red(h):
    return h in ("FF0000",)


def is_orange(h):
    if not h:
        return False
    r, g, b = (int(h[i:i + 2], 16) / 255 for i in (0, 2, 4))
    hh, ss, vv = colorsys.rgb_to_hsv(r, g, b)
    return 15 / 360 <= hh <= 45 / 360 and ss >= 0.45 and vv >= 0.6


def sheet_rank(title):
    t = title.lower()
    y = int(re.search(r"(20\d\d)", t).group(1)) if re.search(r"20\d\d", t) else 2026
    m = max([v for k, v in MONTHS.items() if k in t] or [0])
    return (y, m)


def norm(s):
    return " ".join(re.sub(r"[^a-z ]", " ", (s or "").lower()).split())


def blocks(ws):
    reds = [r for r in range(1, ws.max_row + 1) if sum(is_red(rgb(c)) for c in ws[r][:14]) >= 6]
    edges = [0] + reds + [ws.max_row + 1]
    for a, b in zip(edges, edges[1:]):
        if b - a > 1:
            yield a + 1, b - 1


def main():
    if len(sys.argv) < 3 or sys.argv[1] in ("-h", "--help"):  # OVERLAY_6 4g
        print(__doc__); sys.exit(2)
    wb = load(sys.argv[1])
    args = sys.argv[2:]
    pats = []
    if args and args[0] == "--list":
        for line in open(args[1], encoding="utf-8"):
            if line.strip():
                mrn, _, name = line.rstrip("\n").partition("\t")
                pats.append((mrn.strip(), name.strip()))
    else:
        pats = [(a, "") if a.isdigit() else ("", a) for a in args]
    sheets = sorted([ws for ws in wb.worksheets if "main sheet" in ws.title.lower()],
                    key=lambda w: sheet_rank(w.title), reverse=True)
    index = []  # (sheet, [(a, b, {mrns}, {names})]) built once
    for ws in sheets:
        bl = []
        rows = list(ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=14))
        reds = [i + 1 for i, row in enumerate(rows) if sum(is_red(rgb(c)) for c in row) >= 6]
        edges = [0] + reds + [len(rows) + 1]
        for a, b in zip(edges, edges[1:]):
            if b - a <= 1:
                continue
            ms, ns = set(), set()
            for r in range(a + 1, b):
                row = rows[r - 1]
                vm = str(row[5].value or "").split(".")[0].strip() if len(row) > 5 else ""
                vn = norm(str(row[4].value or "")) if len(row) > 4 else ""
                if vm: ms.add(vm)
                if vn: ns.add(vn)
            bl.append((a + 1, b - 1, ms, ns))
        index.append((ws, rows, bl))
    for mrn, name in pats:
        hit = None
        for ws, rows, bl in index:
            for a, b, ms, ns in bl:
                if (mrn and mrn in ms) or (name and norm(name) in ns and (not mrn or not ms or mrn in ms)):
                    hit = (ws, rows, a, b)
            if hit:
                break
        if not hit:
            print("NOT FOUND\t%s\t%s" % (mrn, name)); continue
        ws, rows, a, b = hit
        # Old admissions may be stacked above the current entry inside the same block:
        # judge only the last entry (from the last row carrying a name or MRN).
        start = a
        for r in range(a, b + 1):
            row = rows[r - 1]
            if len(row) > 5 and (row[4].value or row[5].value):
                start = r
        a = start
        cells = [rgb(c) for r in range(a, b + 1) for c in rows[r - 1][:14]]
        share = sum(is_orange(h) for h in cells) / max(1, len(cells))
        print("%s\t%s\t%s\t%s rows %d-%d orange %.0f%%" % ("FINALISED" if share >= 0.7 else "OPEN",
              mrn, name, ws.title, a, b, share * 100))


if __name__ == "__main__":
    main()
