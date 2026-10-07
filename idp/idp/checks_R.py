#!/usr/bin/env python3
"""Checks R (Chris, 07.10.2026). Run before ANY send or Drive write; every FAIL blocks it.
  checks_R.py text <file> [...]        no doses, no em/en dashes (handout.html, recap_email.txt, delta2.*, id_list_email.txt)
  checks_R.py idlist <id_list_email.txt>  ID list line format
  checks_R.py ams <AMS_rows_*.xlsx>    AMS rows layout of every changed (green) cell
Prints CHECKS R OK or one FAIL line per problem (exit 1)."""
import sys, re
DOSE = re.compile(r"\b\d+(?:[.,]\d+)?\s?(?:mg|g|gr|mcg|µg|ug|IU|MU|units?|mL/kg|mg/kg)\b(?!/dL)|\b\d{3,4}\s?(?:x1|once|OD|BID|TID|q\d+h)\b|\bloading dose of\s*\d|\(\s*\d{3,4}\b(?![,.]\d)|\b\d{3,4}\s+(?:proph|from|then|load|daily|once)\b|\bthen \d{2,4}\b(?!\s*(?:col|/))", re.I)
DASH = re.compile("[–—]")
fails = []
warns = []
def text(paths):
    for p in paths:
        t = open(p, encoding="utf-8").read()
        t2 = re.sub(r"<[^>]+>", " ", t) if p.endswith(".html") else t
        for m in DOSE.finditer(t2):
            ctx = t2[max(0, m.start()-40):m.end()+20].replace("\n", " ")
            if re.search(r"\b(Hb|Hgb|CRP|PCT|Cr|WBC|Plt|lactate|glucose|protein|bili|Na|K|EF|HGT|ANC|col/ml|CrCl)\b", ctx[:45], re.I):
                continue
            fails.append("%s: possible dose '%s' in ...%s..." % (p, m.group(0), ctx))
        for m in DASH.finditer(t):
            fails.append("%s: em/en dash at %d" % (p, m.start()))
LINE = re.compile(r"^(\*[A-Z0-9]+\* [A-Z][a-z'\-]+(?: [A-Z][A-Za-z'\-]+)+ \(new\)|\*[A-Z0-9]+ .+ \(UC\)\*|[A-Z0-9]+ [A-Z][A-Za-z0-9'\-]+(?: [A-Z][A-Za-z'\-()]+)+(?: \(new\))?) (?:[A-Z]{2}|\?)$")
def idlist(p):
    for ln in open(p, encoding="utf-8").read().splitlines():
        if not ln.strip() or ln.endswith(":"):
            continue
        if not LINE.match(ln):
            fails.append("ID list line format: %r" % ln)
        if re.search(r"new on AS list|\(DR\.?\)|[A-Z]{3,} [A-Z]{3,}", ln):
            fails.append("ID list line not in email style (First Last, no AS-list tags): %r" % ln)
        if ln.endswith(" ?"):
            warns.append("ID list initials '?': take them from the email, else yesterday's list, else the AMS last entry; report if still unknown: %r" % ln)
def ams(p):
    import openpyxl
    ws = openpyxl.load_workbook(p)["AMS rows"]
    G = "FFD9F2D0"
    col = lambda r, c: ws.cell(r, c)
    seen = {}
    block = 0
    for r in range(1, ws.max_row + 1):
        a = col(r, 1).value
        if isinstance(a, str) and "|" in a:
            block += 1
        for c in range(1, 35):
            x = col(r, c)
            if x.value is None or x.fill.fgColor.rgb != G or not isinstance(x.value, str):
                continue
            v = x.value
            L = x.column_letter
            if DASH.search(v): fails.append("AMS %s%d: em/en dash" % (L, r))
            if DOSE.search(v): fails.append("AMS %s%d: dose %r" % (L, r, v[:60]))
            if L in ("D", "W", "X") and v not in ("Ongoing", "-") and not re.fullmatch(r"\d{2}/\d{2}/\d{2}", v):
                fails.append("AMS %s%d: date must be dd/mm/yy (or Ongoing): %r" % (L, r, v))
            if L == "AD":
                if not re.match(r"^\d{2}/\d{2}/\d{2} \S", v):
                    fails.append("AMS AD%d: must start 'dd/mm/yy Specimen': %r" % (r, v[:60]))
                for cc, nm in ((31, "AE status"), (32, "AF organism/result"), (33, "AG antibiogram or -")):
                    if not col(r, cc).value:
                        fails.append("AMS AD%d: %s empty; split the culture across AD/AE/AF/AG" % (r, nm))
                if re.search(r"\b(UCx|BCx|cx|neg|pos)\b", v):
                    fails.append("AMS AD%d: abbreviations/result inside AD; specimen only there: %r" % (r, v[:60]))
                k = (block, v[:8], " ".join(re.sub(r"\W+", " ", v[9:].lower()).split()[:3]))
                if k in seen: fails.append("AMS AD%d: duplicate of AD%d" % (r, seen[k]))
                seen[k] = r
            if L == "AE" and v not in ("Final", "Preliminary", "Pending") and not v.startswith("Final (") and not v.startswith("Preliminary ("):
                fails.append("AMS AE%d: status must be Final/Preliminary/Pending: %r" % (r, v))
            if L == "AG" and v != "-" and not re.match(r"^(R: .+; S: .+|S: .+|R: .+|AST .+|Multisensitive.*)$", v):
                fails.append("AMS AG%d: antibiogram must read 'R: x, y; S: a, b' or '-': %r" % (r, v))
            if L == "V" and re.search(r"\b(Mero|Pip|Piptazo|Vanco|Ceft|Teico|Fluco|Fosfo|Erta|Augmentin|Amik|Cefaz|Tava|Azi|Valtrex|Valaci|Vori)\b", v):
                fails.append("AMS V%d: write the full generic name: %r" % (r, v))
            if L == "U" and re.search(r"\(\d{2}/\d{2}\)\s*$", v):
                fails.append("AMS U%d: indication is a diagnosis sentence, no trailing (dd/mm): %r" % (r, v))
            if L == "C" and (re.search(r"\b(Mero|Vanco|Fluco|Cefaz|Fosfo|Erta|Piptazo|Pip|desat|w/|o/n)\b", v) or len(v) < 25):
                fails.append("AMS C%d: plan must be a full-word ID plan like the existing rows: %r" % (r, v[:70]))
mode = sys.argv[1]
if mode == "text": text(sys.argv[2:])
elif mode == "idlist": idlist(sys.argv[2])
elif mode == "ams": ams(sys.argv[2])
for w in warns: print("WARN", w)
for f in fails: print("FAIL", f)
print("CHECKS R OK" if not fails else "CHECKS R FAILED (%d)" % len(fails))
sys.exit(1 if fails else 0)
