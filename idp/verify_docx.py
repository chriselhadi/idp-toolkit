#!/usr/bin/env python3
"""Verify a built handout DOCX against render.json. Prints RESULT: ALL CHECKS PASSED or RESULT: FAILED.

Usage: verify_docx.py <file.docx> <render.json> [--groups "7th floor,ICU#B"]
"""
import sys, os, zipfile, argparse, re
sys.path.insert(0, os.path.dirname(__file__))
from common import WARD_ORDER, jload
from docx import Document
from docx.oxml.ns import qn


def main():
    ap = argparse.ArgumentParser(); ap.add_argument("docx"); ap.add_argument("render"); ap.add_argument("--groups")
    a = ap.parse_args()
    groups = [g.strip() for g in a.groups.split(",")] if a.groups else None
    render = jload(a.render)
    expected = [p for p in render["patients"] if not groups or p["group"] in groups]
    fails = []
    try:
        z = zipfile.ZipFile(a.docx); bad = z.testzip()
        if bad: fails.append("zip member corrupt: %s" % bad)
        if "word/document.xml" not in z.namelist(): fails.append("no word/document.xml")
    except Exception as e:
        print("RESULT: FAILED (not a zip: %s)" % e); return 1
    doc = Document(a.docx)
    if not doc.tables: print("RESULT: FAILED (no table)"); return 1
    table = doc.tables[0]
    rows = table.rows
    pat_rows = []
    order_seen = []
    for row in rows[1:]:
        cells = row.cells
        texts = [c.text for c in cells]
        if len(set(id(c._tc) for c in cells)) == 1:  # merged group heading
            order_seen.append(texts[0].strip()); continue
        pat_rows.append((row, texts))
    if len(pat_rows) != len(expected):
        fails.append("patient rows %d != expected %d" % (len(pat_rows), len(expected)))
    for (row, texts), p in zip(pat_rows, expected):
        c1 = texts[0]
        if p["name"] not in c1: fails.append("row for %s does not carry the name" % p["pid"])
        for line in p["lines_docx"]:
            if line.strip() and re.sub(r"\s+", " ", line.strip()) not in re.sub(r"\s+", " ", c1):
                fails.append("%s: line missing from column 1: %s" % (p["pid"], line[:60]))
        if texts[1].strip() or texts[2].strip():
            fails.append("%s: columns 2/3 are not empty" % p["pid"])
        trPr = row._tr.trPr
        h = trPr.find(qn("w:trHeight")) if trPr is not None else None
        if h is None or int(h.get(qn("w:val"), "0")) < 1100:
            fails.append("%s: row height below 1100 twips" % p["pid"])
        if "—" in c1 or "–" in c1: fails.append("%s: dash in column 1" % p["pid"])
    exp_groups = [g for g in WARD_ORDER if any(p["group"] == g for p in expected)]
    if order_seen != exp_groups: fails.append("ward order %s != %s" % (order_seen, exp_groups))
    full = "\n".join(t[0] for _, t in pat_rows)
    if len(full.strip()) < 50 * max(1, len(expected)): fails.append("column 1 nearly empty")
    for f in fails: print("FAIL", f)
    print("checked %d patient rows, %d groups" % (len(pat_rows), len(order_seen)))
    print("RESULT: ALL CHECKS PASSED" if not fails else "RESULT: FAILED")
    return 1 if fails else 0


if __name__ == "__main__":
    sys.exit(main())
