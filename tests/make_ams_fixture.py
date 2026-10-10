#!/usr/bin/env python3
"""A small fictional AMS workbook in the real sheet's layout, for ams_check.py in the end-to-end run.

    python3 make_ams_fixture.py <out.xlsx>

One tab "Main sheet Sep 2026": header row, red separator rows between blocks, columns A:AH as in ams_rows.py.
  Alpha Tester One   (MRN 100001)  Mero ongoing (agrees), Vanco left "Ongoing" although the record gave a single
                                    dose (X_AMS_COURSE_OPEN), urine culture 04/09 (agrees), blood culture 05/09 with
                                    K. pneumoniae that the record lacks (X_AMS_CULTURE_NOT_ON_LIST)
  Bravo Tester Two   (MRN 100002)  entry shaded orange as finalised while still on the list (X_AMS_FINALISED_BUT_ACTIVE)
  Charlie Tester Three (100003)    no antimicrobial row while Erta runs (X_COURSE_NOT_IN_AMS)
Every other census patient has no entry (X_NOT_IN_AMS).
"""
import datetime as dt
import sys

import openpyxl
from openpyxl.styles import PatternFill

HEAD = ["", "FELLOW", "Plan", "Plan Date", "Pt name", "MRN", "Age", "Room", "weight", "Creatinine",
        "CrCl", "Date of Admission", "Date of ID Consultation", "Chief complaint", "PMH", "Active Issues",
        "last hospitalization", "Abx history", "Imaging", "Pertinent labs", "Indication/ID Diagnosis",
        "Actual antibiotics", "Date started", "Date ended", "Released by ID", "Duration advice",
        "Attending", "Lines", "Discharge Abx", "Culture results/Notes", "Preliminary Results", "Organism",
        "Antibiogram", "notes"]
RED = PatternFill("solid", fgColor="FFFF0000")
ORANGE = PatternFill("solid", fgColor="FFF6B26B")
D = lambda d, m: dt.datetime(2026, m, d)


def main(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Main sheet Sep 2026"
    ws.append(HEAD)

    def red():
        ws.append([""] * 34)
        for c in range(1, 15):
            ws.cell(ws.max_row, c).fill = RED

    def row(**cells):
        r = [""] * 34
        for col, v in cells.items():
            r[int(col[1:]) - 1] = v
        ws.append(r)
        return ws.max_row

    red()
    row(c2="CH", c5="Alpha Tester One", c6="100001", c7="70", c8="708B", c12=D(27, 8), c21="Urosepsis",
        c22="Meropenem", c23=D(3, 9), c24="Ongoing", c30="04/09/26 Urine", c31="Final", c32="E. coli ESBL",
        c33="R: Ceftriaxone; S: Meropenem")
    row(c22="Vancomycin", c23=D(3, 9), c24="Ongoing", c30="05/09/26 Blood", c31="Final", c32="Klebsiella pneumoniae",
        c33="S: Meropenem")
    red()
    first = row(c2="MA", c5="Bravo Tester Two", c6="100002", c8="ICU4B", c21="HAP", c22="Piperacillin-tazobactam",
                c23=D(29, 8), c24="Ongoing")
    for c in range(1, 35):
        ws.cell(first, c).fill = ORANGE
    red()
    row(c2="CH", c5="Charlie Tester Three", c6="100003", c8="425", c21="Diabetic foot")
    red()
    wb.save(path)
    print("AMS fixture: %s" % path)


if __name__ == "__main__":
    main(sys.argv[1])
